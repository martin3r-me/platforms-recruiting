"""
Kundenfolien "Schulungstermine umplanen und nachbereiten" im Stil der
RheinGedeck-HCM-Folien (blauer Rand links, Kicker, Titel, Schritte links,
Screenshot rechts, Fusszeile "RheinGedeck · Intern").

Aufruf:
    python schulungstermine_folien.py <screenshot-ordner> <ziel.pptx>

Screenshots liegen anonymisiert in folien/screenshots/ (nie echte Namen/Fotos). Erwartete Dateinamen
siehe SCREENSHOTS unten; fehlt eine Datei, steht dort ein beschrifteter
Platzhalter. Braucht python-pptx und Pillow.
"""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BLAU = RGBColor(0x12, 0x73, 0xDE)
SCHWARZ = RGBColor(0x11, 0x11, 0x11)
GRAU = RGBColor(0x5F, 0x63, 0x68)
LINIE = RGBColor(0xB8, 0xBC, 0xC4)
PLATZHALTER = RGBColor(0xF1, 0xF3, 0xF5)
FONT = "Calibri"

# Dateiname → was darauf zu sehen sein soll (steht auch im Platzhalter).
SCREENSHOTS = {
    "A_liste_angehakt.png": "Teilnehmerliste (Übersicht) mit angehakten Personen und blauem Balken „N ausgewählt · Verschieben nach…“",
    "B_fenster_verschieben.png": "Fenster „Teilnehmer verschieben“ mit Zieltermin-Auswahl und Kommentarfeld",
    "C_status_teilgenommen.png": "Eine Zeile der Teilnehmerliste mit Status-Auswahl „Teilgenommen“",
    "D_nach_der_schulung.png": "Reiter „Nach der Schulung“: Spalten Zuschlag, Vertragslaufzeit, Versand",
    "E_statusfilter.png": "Statusfilter aufgeklappt, Eintrag „Keine Reaktion“ sichtbar",
    "F_statistik_pille.png": "Statistik → Schulungstermine: Badge „Ausgebucht“ und Pille „N ohne Termin“",
    "G_kampagne_fenster.png": "Kampagnen-Fenster: Anlass-Karte, Empfängerliste, Nachricht, Senden-Knopf",
}


def text(slide, x, y, w, h, inhalt, groesse, fett=False, farbe=SCHWARZ, ausrichtung=PP_ALIGN.LEFT, anker=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anker
    tf.margin_left = tf.margin_right = Inches(0.05)
    zeilen = inhalt if isinstance(inhalt, list) else [inhalt]
    for i, zeile in enumerate(zeilen):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ausrichtung
        p.space_after = Pt(groesse * 0.6)
        # **fett** innerhalb einer Zeile
        teile = zeile.split("**")
        for j, teil in enumerate(teile):
            if not teil:
                continue
            r = p.add_run()
            r.text = teil
            r.font.name = FONT
            r.font.size = Pt(groesse)
            r.font.bold = fett or (j % 2 == 1)
            r.font.color.rgb = farbe
    return box


def rahmen(slide, nr, kicker, titel):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.17), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLAU; bar.line.fill.background()
    text(slide, 0.67, 0.36, 11.46, 0.3, kicker.upper(), 12, fett=True, farbe=BLAU)
    text(slide, 0.67, 0.80, 11.96, 0.7, titel, 28.5, fett=True)
    linie = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.56), Inches(1.57), Inches(12.17), Emu(9144))
    linie.fill.solid(); linie.fill.fore_color.rgb = LINIE; linie.line.fill.background()
    text(slide, 0.67, 7.02, 6.77, 0.3, "RheinGedeck · Intern", 10.5, farbe=GRAU)
    text(slide, 12.04, 7.02, 0.58, 0.3, f"{nr:02d}", 10.5, farbe=GRAU, ausrichtung=PP_ALIGN.RIGHT)


def bild(slide, ordner, datei, x, y, w, h):
    """Screenshot einpassen (Seitenverhaeltnis bleibt), sonst Platzhalter."""
    pfad = ordner / datei
    if pfad.exists():
        with Image.open(pfad) as im:
            bw, bh = im.size
        faktor = min(w / bw, h / bh)
        pw, ph = bw * faktor, bh * faktor
        pic = slide.shapes.add_picture(str(pfad), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        pic.line.color.rgb = LINIE
        pic.line.width = Pt(0.75)
        return
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.04
    box.fill.solid(); box.fill.fore_color.rgb = PLATZHALTER
    box.line.color.rgb = LINIE; box.line.width = Pt(1.25); box.line.dash_style = 4  # gestrichelt
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (zeile, groesse, fett) in enumerate([(f"Screenshot {datei.split('_')[0]}", 16, True), (SCREENSHOTS[datei], 13, False), (datei, 10, False)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = zeile
        r.font.name = FONT; r.font.size = Pt(groesse); r.font.bold = fett; r.font.color.rgb = GRAU


def hinweis(slide, x, y, w, zeilen):
    """Blaue Kernaussagen rechts neben einem Bild (wie Folie 2 im Original)."""
    text(slide, x, y, w, 4, zeilen, 19, fett=True, farbe=BLAU)


def bauen(ordner: Path, ziel: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    leer = prs.slide_layouts[6]
    nr = 0

    def folie(kicker, titel):
        nonlocal nr
        nr += 1
        s = prs.slides.add_slide(leer)
        rahmen(s, nr, kicker, titel)
        return s

    # 1 — verschieben: auswaehlen
    s = folie("Schritt 1 · Teilnehmer verschieben", "Teilnehmer auswählen")
    text(s, 0.75, 1.94, 3.6, 4.2, [
        "1. **Schulungstermine** öffnen und den Termin anklicken, aus dem ihr verschieben wollt.",
        "2. In der Teilnehmerliste die Personen **anhaken**. Das Kästchen in der Kopfzeile wählt alle auf einmal.",
        "3. Im blauen Balken auf **„Verschieben nach…“** klicken.",
    ], 17)
    text(s, 0.75, 5.9, 3.6, 0.9, "Geht vorher und während der Schulung.", 17, fett=True, farbe=BLAU)
    bild(s, ordner, "A_liste_angehakt.png", 4.55, 1.88, 8.05, 5.0)

    # 2 — verschieben: Ziel
    s = folie("Schritt 2 · Teilnehmer verschieben", "Zieltermin wählen")
    text(s, 0.75, 1.94, 3.6, 4.2, [
        "4. Den **Zieltermin** wählen. Hinter jedem Termin steht, wie viele Plätze frei sind.",
        "5. Optional einen **Kommentar** eintragen, z. B. „Logistik-Gruppe“. Er steht im Verlauf.",
        "6. Auf **„Verschieben“** klicken.",
    ], 17)
    text(s, 0.75, 5.7, 3.6, 1.1, "Status, Bestätigung und Notizen bleiben. Es geht keine Nachricht raus.", 17, fett=True, farbe=BLAU)
    bild(s, ordner, "B_fenster_verschieben.png", 4.55, 1.88, 8.05, 5.0)

    # 3 — Schulungsabend
    s = folie("Am Schulungsabend", "Erst „Teilgenommen“, dann aufteilen")
    text(s, 0.75, 1.91, 11.8, 0.8, "In der großen Gruppe alle Anwesenden auf **Teilgenommen** setzen. Danach in die Gruppen verschieben. Der Status wandert mit.", 20)
    bild(s, ordner, "C_status_teilgenommen.png", 0.75, 2.86, 7.97, 3.9)
    hinweis(s, 9.18, 3.0, 3.4, ["Ziel: gleiche Stelle", "Bis der Zieltermin zu Ende ist", "Teilgenommen nur am selben Tag"])
    text(s, 9.18, 5.75, 3.4, 1.0, "Volle Gruppe: Wer nicht mehr passt, bleibt im alten Termin.", 16)

    # 4 — Nachbereitung
    s = folie("Nach der Schulung", "Vertragsdaten eintragen, auch zu zweit")
    text(s, 0.75, 1.94, 3.6, 4.2, [
        "1. Den Termin öffnen, Reiter **„Nach der Schulung“**.",
        "2. **Zuschlag**, **Vertragsbeginn** und bei Bedarf **Vertragsende** eintragen. Ohne Ende rechnet das System es aus.",
        "3. Zum Schluss die Verträge über den **Sammelversand** schicken.",
    ], 17)
    text(s, 0.75, 5.6, 3.6, 1.2, "Jede Eingabe ist sofort gespeichert. Zwei Laptops sehen denselben Stand.", 17, fett=True, farbe=BLAU)
    bild(s, ordner, "D_nach_der_schulung.png", 4.55, 1.88, 8.05, 5.0)

    # 5 — Filter
    s = folie("Überblick behalten", "Nach Status filtern, auch „Keine Reaktion“")
    text(s, 0.75, 1.94, 3.6, 4.2, [
        "Oben in der Liste das Feld **Status** öffnen. Der Filter gilt in der Übersicht und nach der Schulung.",
        "**Keine Reaktion** zeigt Gebuchte, die auf die Erinnerungen nicht reagiert haben. Ihr Platz ist wieder frei.",
    ], 17)
    text(s, 0.75, 5.9, 3.6, 0.9, "Praktisch zum Nachtelefonieren vor der Schulung.", 17, fett=True, farbe=BLAU)
    bild(s, ordner, "E_statusfilter.png", 4.55, 1.88, 8.05, 5.0)

    # 6/7 — Kampagne. Ohne beide Bilder eine Textfolie in zwei Spalten statt
    # zweier Folien mit Platzhaltern (Bilder gibt es nur, wenn gerade ein
    # kommender Termin voll ist).
    kampagne_bilder = (ordner / "F_statistik_pille.png").exists() or (ordner / "G_kampagne_fenster.png").exists()
    if kampagne_bilder:
        s = folie("Schulung voll · Schritt 1", "Bewerber ohne Termin finden")
        text(s, 0.75, 1.94, 3.6, 4.2, [
            "1. **Statistik** öffnen, Filiale wählen, zu **„Schulungstermine“** scrollen.",
            "2. Am vollen Termin auf die Pille **„N ohne Termin“** klicken.",
        ], 17)
        text(s, 0.75, 5.6, 3.6, 1.2, "Die Pille erscheint nur bei voller Platzzahl, künftigem Termin und hinterlegter Ausschreibung.", 16, farbe=GRAU)
        bild(s, ordner, "F_statistik_pille.png", 4.55, 1.88, 8.05, 5.0)

        s = folie("Schulung voll · Schritt 2", "Auf freie Termine hinweisen")
        text(s, 0.75, 1.94, 3.6, 4.4, [
            "3. Oben die **Anlass-Karte** lesen. Ist sie rot, gibt es keinen freien Termin: erst einen anlegen.",
            "4. Empfänger prüfen, wer nicht soll: Haken raus.",
            "5. Auf **„WhatsApp an N Personen senden“** klicken und bestätigen.",
        ], 17)
        bild(s, ordner, "G_kampagne_fenster.png", 4.55, 1.88, 8.05, 5.0)
    else:
        s = folie("Wenn eine Schulung voll ist", "Bewerber ohne Termin auf freie Termine hinweisen")
        text(s, 0.75, 1.94, 5.6, 4.4, [
            "1. **Statistik** öffnen, Filiale wählen, zu **„Schulungstermine“** scrollen.",
            "2. Am vollen Termin auf die Pille **„N ohne Termin“** klicken.",
            "3. Oben die **Anlass-Karte** lesen. Ist sie rot, gibt es keinen freien Termin: erst einen anlegen.",
        ], 19.5)
        text(s, 6.9, 1.94, 5.6, 3.2, [
            "4. Empfänger prüfen, wer nicht soll: Haken raus.",
            "5. Auf **„WhatsApp an N Personen senden“** klicken und bestätigen.",
        ], 19.5)
        hinweis(s, 6.9, 4.6, 5.6, ["Die Bewerber buchen sich selbst in einen freien Termin."])
        text(s, 0.75, 6.1, 11.8, 0.7, "Die Pille erscheint nur bei voller Platzzahl, künftigem Termin und hinterlegter Ausschreibung.", 16, farbe=GRAU)

    # 8 — Gut zu wissen
    s = folie("Gut zu wissen", "Regeln auf einen Blick")
    zeilen = [
        ("Verschieben: Termine", "Nur aktive Termine derselben Stelle, die noch nicht zu Ende sind."),
        ("Verschieben: wer", "Gebucht, registriert, bestätigt immer. Teilgenommen nur am selben Tag."),
        ("Verschieben: Nachricht", "Keine. Soll die Person Bescheid wissen, bitte selbst anschreiben."),
        ("Gruppengröße", "Für kleinere Gruppen beim Termin eine kleinere Platzzahl eintragen."),
        ("Vertragsbeginn geleert", "Eine offene Versand-Vormerkung wird zurückgenommen."),
        ("Nach dem Versand", "Es gelten die Daten im Vertrag. Änderung über „Vertrag neu ausstellen“."),
        ("Kampagne: doppelt", "Wer in 14 Tagen schon eine Kampagne bekam, ist vorsichtshalber abgehakt."),
    ]
    tab = s.shapes.add_table(len(zeilen), 2, Inches(0.75), Inches(1.95), Inches(11.8), Inches(4.8)).table
    tab.columns[0].width = Inches(3.4); tab.columns[1].width = Inches(8.4)
    tab.first_row = False
    for i, (a, b) in enumerate(zeilen):
        for j, wert in enumerate((a, b)):
            zelle = tab.cell(i, j)
            zelle.fill.solid(); zelle.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else PLATZHALTER
            zelle.margin_left = Inches(0.12)
            p = zelle.text_frame.paragraphs[0]
            r = p.add_run(); r.text = wert
            r.font.name = FONT; r.font.size = Pt(16); r.font.bold = j == 0
            r.font.color.rgb = BLAU if j == 0 else SCHWARZ

    prs.save(ziel)
    fehlend = [d for d in SCREENSHOTS if not (ordner / d).exists()
               and (kampagne_bilder or not d.startswith(("F_", "G_")))]
    print(f"{ziel} gespeichert, {nr} Folien.")
    if fehlend:
        print("Platzhalter für:", ", ".join(fehlend))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    bauen(Path(sys.argv[1]).expanduser(), Path(sys.argv[2]).expanduser())
